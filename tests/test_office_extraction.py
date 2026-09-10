"""Spreadsheets and decks — the two formats the upload door claimed and never had.

`documents/router.py` listed the xlsx and pptx MIME types in a `_NATIVE_MIMES` set, which read
like support and was referenced by nothing; `extract_native_text` had no branch for either and no
parser existed in the tree. An uploaded Excel therefore reached `route_document` with no text
layer, no page image and no audio, and came back `unsupported` — the dashboard rendered "No
extractable text found in this file" for a file whose whole content is text.

These tests build real workbooks and real decks rather than fixtures, so they fail if the parser
regresses AND if the library is missing from requirements.txt.
"""
from __future__ import annotations

import io

import pytest

from genios_engine.capture.documents import native
from genios_engine.capture.documents.native import extract_text_best_effort

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _workbook() -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Q3 Pipeline"
    ws.append(["Company", "Owner", "Stage", "Value"])
    ws.append(["Antler", "Rohit", "Proposal, sent", 2500000])
    ws.append([None, None, None, None])                  # blank row → dropped, not an empty line
    ws.append(["Inkbox", "Pratap", "Closed Won", 300000])
    wb.create_sheet("Notes").append(["Renewal due 15 Oct"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _deck() -> bytes:
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Pricing"
    slide.placeholders[1].text = "Startup plan: Rs 25,000/mo"
    slide.notes_slide.notes_text_frame.text = "Rohit pushed back on this number."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_xlsx_every_sheet_and_cell_is_read():
    out = extract_text_best_effort(mime=XLSX_MIME, data=_workbook(), filename="pipeline.xlsx")
    for token in ("Antler", "Rohit", "2500000", "Inkbox", "Closed Won", "Renewal due 15 Oct"):
        assert token in out, f"{token!r} lost on the way out of the workbook"


def test_xlsx_sheet_names_become_markdown_headings():
    """Not cosmetic: `chunking.SECTION` splits on markdown headings, so this is the difference
    between an evidence span that cites *Q3 Pipeline* and one that cites a character offset."""
    out = extract_text_best_effort(mime=XLSX_MIME, data=_workbook(), filename="pipeline.xlsx")
    assert "## Sheet: Q3 Pipeline" in out and "## Sheet: Notes" in out


def test_xlsx_rows_are_tab_separated_so_commas_inside_cells_survive():
    """"Proposal, sent" is one cell. Joining rows with commas would silently turn a four-column
    sheet into a five-column one, and every column after it would name the wrong field."""
    out = extract_text_best_effort(mime=XLSX_MIME, data=_workbook(), filename="pipeline.xlsx")
    row = next(ln for ln in out.splitlines() if ln.startswith("Antler"))
    assert row.split("\t") == ["Antler", "Rohit", "Proposal, sent", "2500000"]


def test_xlsx_is_read_by_extension_even_when_the_browser_mislabels_the_mime():
    """Older exporters and some browsers send application/vnd.ms-excel for a real .xlsx."""
    out = extract_text_best_effort(mime="application/vnd.ms-excel", data=_workbook(),
                                   filename="pipeline.xlsx")
    assert "Antler" in out


def test_legacy_xls_is_routed_to_xlrd_not_declared_unsupported():
    """.xls is a different binary format that openpyxl cannot open at all. Asserted at the
    dispatch, because building a real .xls needs a writer library the product does not ship."""
    seen: dict[str, bytes] = {}

    def fake(raw: bytes) -> str:
        seen["raw"] = raw
        return "## Sheet: Sheet1\nold format"

    original = native._xls_to_text
    native._xls_to_text = fake
    try:
        out = extract_text_best_effort(mime="application/vnd.ms-excel", data=b"\xd0\xcf\x11\xe0",
                                       filename="legacy.xls")
    finally:
        native._xls_to_text = original
    assert seen["raw"] == b"\xd0\xcf\x11\xe0" and "old format" in out


def test_pptx_reads_slides_and_speaker_notes():
    """The slide shows the number; the note says why it moved. Dropping notes would discard the
    half a reader actually needs."""
    out = extract_text_best_effort(mime=PPTX_MIME, data=_deck(), filename="deck.pptx")
    assert "Startup plan: Rs 25,000/mo" in out
    assert "Rohit pushed back on this number." in out


def test_pptx_slides_become_markdown_headings():
    out = extract_text_best_effort(mime=PPTX_MIME, data=_deck(), filename="deck.pptx")
    assert out.startswith("## Slide 1: Pricing")


@pytest.mark.parametrize("name, mime", [
    ("pipeline.xlsx", XLSX_MIME),
    ("deck.pptx", PPTX_MIME),
])
def test_a_corrupt_office_file_returns_empty_not_garbage(name, mime):
    """A truncated upload must read as "no extractable text", never as decoded binary — the
    parsers raise, `extract_native_text` swallows, and the plaintext fallback does not apply."""
    assert extract_text_best_effort(mime=mime, data=b"PK\x03\x04garbage", filename=name) == ""
